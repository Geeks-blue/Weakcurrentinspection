from datetime import datetime
from io import BytesIO, StringIO

import qrcode
from fastapi import APIRouter, Body, Depends, File, HTTPException, Response, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import Column, DateTime, Integer, String
from sqlalchemy.orm import Session, declarative_base

from app.deps import get_db, require_teacher_or_admin
from app.models.entities import Asset, Building, Room, User
from app.schemas.assets import AssetItemView, AssetRoomItem, ImportSummary

router = APIRouter()
BaseTmp = declarative_base()

class RoomSignLog(BaseTmp):
    __tablename__ = "room_sign_log"
    id = Column(Integer, primary_key=True)
    room_id = Column(Integer, nullable=False)
    user_id = Column(Integer, nullable=True)
    sign_time = Column(DateTime, default=datetime.utcnow)
    sign_type = Column(String(32), default="scan")
    extra = Column(String(255), nullable=True)

class AssetSignLog(BaseTmp):
    __tablename__ = "asset_sign_log"
    id = Column(Integer, primary_key=True)
    asset_id = Column(Integer, nullable=False)
    user_id = Column(Integer, nullable=True)
    sign_time = Column(DateTime, default=datetime.utcnow)
    sign_type = Column(String(32), default="scan")
    extra = Column(String(255), nullable=True)
# 房间扫码签到接口
@router.post("/room/{room_id}/sign")
def sign_room(room_id: int, db: Session = Depends(get_db), user: User = Depends(require_teacher_or_admin), sign_type: str = Body("scan"), extra: str = Body(None)):
    room = db.query(Room).filter(Room.id == room_id).first()
    if not room:
        raise HTTPException(status_code=404, detail="房间不存在")
    log = RoomSignLog(room_id=room_id, user_id=user.id, sign_time=datetime.utcnow(), sign_type=sign_type, extra=extra)
    db.add(log)
    db.commit()
    return {"ok": True, "sign_time": log.sign_time}

# 资产扫码签到接口
@router.post("/item/{asset_id}/sign")
def sign_asset(asset_id: int, db: Session = Depends(get_db), user: User = Depends(require_teacher_or_admin), sign_type: str = Body("scan"), extra: str = Body(None)):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="资产不存在")
    log = AssetSignLog(asset_id=asset_id, user_id=user.id, sign_time=datetime.utcnow(), sign_type=sign_type, extra=extra)
    db.add(log)
    db.commit()
    return {"ok": True, "sign_time": log.sign_time}
@router.post("/room", response_model=AssetRoomItem)
def create_room(item: AssetRoomItem, db: Session = Depends(get_db), _: User = Depends(require_teacher_or_admin)):
    building = db.query(Building).filter(Building.code == item.building_code).first()
    if not building:
        raise HTTPException(status_code=400, detail="楼栋不存在")
    room = Room(
        room_code=item.room_code,
        building_id=building.id,
        floor_label=item.floor_label,
        location_text=item.location_text,
        qr_token=f"QR-{item.room_code}",
        is_active=item.is_active,
    )
    db.add(room)
    db.commit()
    db.refresh(room)
    return AssetRoomItem(
        room_id=room.id,
        building_code=building.code,
        building_name=building.name,
        room_code=room.room_code,
        floor_label=room.floor_label,
        location_text=room.location_text,
        is_active=room.is_active,
    )

@router.put("/room/{room_id}", response_model=AssetRoomItem)
def update_room(room_id: int, item: AssetRoomItem, db: Session = Depends(get_db), _: User = Depends(require_teacher_or_admin)):
    room = db.query(Room).filter(Room.id == room_id).first()
    if not room:
        raise HTTPException(status_code=404, detail="房间不存在")
    building = db.query(Building).filter(Building.code == item.building_code).first()
    if not building:
        raise HTTPException(status_code=400, detail="楼栋不存在")
    room.building_id = building.id
    room.room_code = item.room_code
    room.qr_token = f"QR-{item.room_code}"
    room.floor_label = item.floor_label
    room.location_text = item.location_text
    room.is_active = item.is_active
    db.commit()
    db.refresh(room)
    return AssetRoomItem(
        room_id=room.id,
        building_code=building.code,
        building_name=building.name,
        room_code=room.room_code,
        floor_label=room.floor_label,
        location_text=room.location_text,
        is_active=room.is_active,
    )

@router.delete("/room/{room_id}")
def delete_room(room_id: int, db: Session = Depends(get_db), _: User = Depends(require_teacher_or_admin)):
    room = db.query(Room).filter(Room.id == room_id).first()
    if not room:
        raise HTTPException(status_code=404, detail="房间不存在")
    db.delete(room)
    db.commit()
    return {"ok": True}

@router.get("/room/{room_id}/qrcode")
def get_room_qrcode(room_id: int, db: Session = Depends(get_db)):
    room = db.query(Room).filter(Room.id == room_id).first()
    if not room:
        raise HTTPException(status_code=404, detail="房间不存在")
    qr = qrcode.QRCode(box_size=8, border=2)
    qr.add_data(room.qr_token)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buf = BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return Response(content=buf.read(), media_type="image/png")
@router.post("/item", response_model=AssetItemView)
def create_asset(item: AssetItemView, db: Session = Depends(get_db), _: User = Depends(require_teacher_or_admin)):
    room = db.query(Room).filter(Room.room_code == item.room_code).first()
    if not room:
        raise HTTPException(status_code=400, detail="房间不存在")
    asset = Asset(
        asset_code=item.asset_code,
        asset_name=item.asset_name,
        asset_category=item.asset_category,
        room_id=room.id,
        quantity=item.quantity,
        status=item.status,
        manufacturer=item.manufacturer,
        model=item.model,
        note=item.note,
    )
    db.add(asset)
    db.commit()
    db.refresh(asset)
    building = db.query(Building).filter(Building.id == room.building_id).first()
    return AssetItemView(
        asset_id=asset.id,
        asset_code=asset.asset_code,
        asset_name=asset.asset_name,
        asset_category=asset.asset_category,
        building_code=building.code,
        room_code=room.room_code,
        quantity=asset.quantity,
        status=asset.status,
        manufacturer=asset.manufacturer,
        model=asset.model,
        note=asset.note,
        updated_at=asset.updated_at,
    )

@router.put("/item/{asset_id}", response_model=AssetItemView)
def update_asset(asset_id: int, item: AssetItemView, db: Session = Depends(get_db), _: User = Depends(require_teacher_or_admin)):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="资产不存在")
    room = db.query(Room).filter(Room.room_code == item.room_code).first()
    if not room:
        raise HTTPException(status_code=400, detail="房间不存在")
    asset.asset_code = item.asset_code
    asset.asset_name = item.asset_name
    asset.asset_category = item.asset_category
    asset.room_id = room.id
    asset.quantity = item.quantity
    asset.status = item.status
    asset.manufacturer = item.manufacturer
    asset.model = item.model
    asset.note = item.note
    db.commit()
    db.refresh(asset)
    building = db.query(Building).filter(Building.id == room.building_id).first()
    return AssetItemView(
        asset_id=asset.id,
        asset_code=asset.asset_code,
        asset_name=asset.asset_name,
        asset_category=asset.asset_category,
        building_code=building.code,
        room_code=room.room_code,
        quantity=asset.quantity,
        status=asset.status,
        manufacturer=asset.manufacturer,
        model=asset.model,
        note=asset.note,
        updated_at=asset.updated_at,
    )

@router.delete("/item/{asset_id}")
def delete_asset(asset_id: int, db: Session = Depends(get_db), _: User = Depends(require_teacher_or_admin)):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="资产不存在")
    db.delete(asset)
    db.commit()
    return {"ok": True}

@router.get("/item/{asset_id}/qrcode")
def get_asset_qrcode(asset_id: int, db: Session = Depends(get_db)):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="资产不存在")
    qr = qrcode.QRCode(box_size=8, border=2)
    qr.add_data(asset.asset_code)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buf = BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return Response(content=buf.read(), media_type="image/png")
def _parse_table_file(file: UploadFile) -> list[dict[str, str]]:
    content = file.file.read()
    if not content:
        return []

    filename = (file.filename or "").lower()
    if filename.endswith(".xlsx"):
        workbook = load_workbook(filename=BytesIO(content), data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(item or "").strip() for item in rows[0]]
        result: list[dict[str, str]] = []
        for row in rows[1:]:
            data: dict[str, str] = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = row[idx] if idx < len(row) else ""
                data[header] = "" if value is None else str(value).strip()
            if any(data.values()):
                result.append(data)
        return result

    if filename.endswith(".csv"):
        text = ""
        for encoding in ("utf-8-sig", "gbk", "latin-1"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if not text:
            return []

        import csv

        reader = csv.DictReader(StringIO(text))
        result = []
        for row in reader:
            normalized = {str(k or "").strip(): str(v or "").strip() for k, v in row.items() if k}
            if any(normalized.values()):
                result.append(normalized)
        return result

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="仅支持 .csv 或 .xlsx 文件",
    )


@router.get("/rooms", response_model=list[AssetRoomItem])
def list_rooms(
    db: Session = Depends(get_db),
    _: User = Depends(require_teacher_or_admin),
) -> list[AssetRoomItem]:
    rows = (
        db.query(Room, Building)
        .join(Building, Room.building_id == Building.id)
        .order_by(Building.code.asc(), Room.room_code.asc())
        .all()
    )
    return [
        AssetRoomItem(
            room_id=room.id,
            building_code=building.code,
            building_name=building.name,
            room_code=room.room_code,
            floor_label=room.floor_label,
            location_text=room.location_text,
            is_active=room.is_active,
        )
        for room, building in rows
    ]


@router.get("/items", response_model=list[AssetItemView])
def list_assets(
    db: Session = Depends(get_db),
    _: User = Depends(require_teacher_or_admin),
) -> list[AssetItemView]:
    rows = (
        db.query(Asset, Room, Building)
        .join(Room, Asset.room_id == Room.id)
        .join(Building, Room.building_id == Building.id)
        .order_by(Asset.updated_at.desc())
        .all()
    )
    return [
        AssetItemView(
            asset_id=asset.id,
            asset_code=asset.asset_code,
            asset_name=asset.asset_name,
            asset_category=asset.asset_category,
            building_code=building.code,
            room_code=room.room_code,
            quantity=int(asset.quantity),
            status=asset.status,
            manufacturer=asset.manufacturer,
            model=asset.model,
            note=asset.note,
            updated_at=asset.updated_at,
        )
        for asset, room, building in rows
    ]


@router.post("/import/rooms", response_model=ImportSummary)
def import_rooms(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: User = Depends(require_teacher_or_admin),
) -> ImportSummary:
    rows = _parse_table_file(file)
    created_count = 0
    updated_count = 0
    skipped_count = 0
    errors: list[str] = []

    for index, row in enumerate(rows, start=2):
        building_code = row.get("building_code", "").strip()
        room_code = row.get("room_code", "").strip()
        if not building_code or not room_code:
            skipped_count += 1
            errors.append(f"第{index}行缺少 building_code 或 room_code")
            continue

        building = db.query(Building).filter(Building.code == building_code).first()
        if not building:
            building = Building(
                code=building_code,
                name=row.get("building_name", building_code),
                category=row.get("building_category", "public") or "public",
                gender_restriction=row.get("gender_restriction", "none") or "none",
            )
            db.add(building)
            db.flush()

        room = db.query(Room).filter(Room.room_code == room_code).first()
        is_active_text = row.get("is_active", "true").lower()
        is_active = is_active_text not in {"0", "false", "no", "n"}

        if room:
            room.building_id = building.id
            room.floor_label = row.get("floor_label") or room.floor_label
            room.location_text = row.get("location_text") or room.location_text
            room.is_active = is_active
            room.qr_token = row.get("qr_token") or f"QR-{room_code}"
            updated_count += 1
        else:
            room = Room(
                room_code=room_code,
                building_id=building.id,
                floor_label=row.get("floor_label") or None,
                location_text=row.get("location_text") or None,
                qr_token=row.get("qr_token") or f"QR-{room_code}",
                is_active=is_active,
            )
            db.add(room)
            created_count += 1

    db.commit()
    return ImportSummary(
        total_rows=len(rows),
        created_count=created_count,
        updated_count=updated_count,
        skipped_count=skipped_count,
        errors=errors,
    )


@router.post("/import/items", response_model=ImportSummary)
def import_assets(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: User = Depends(require_teacher_or_admin),
) -> ImportSummary:
    rows = _parse_table_file(file)
    created_count = 0
    updated_count = 0
    skipped_count = 0
    errors: list[str] = []

    for index, row in enumerate(rows, start=2):
        asset_code = row.get("asset_code", "").strip()
        asset_name = row.get("asset_name", "").strip()
        room_code = row.get("room_code", "").strip()

        if not asset_code or not asset_name or not room_code:
            skipped_count += 1
            errors.append(f"第{index}行缺少 asset_code / asset_name / room_code")
            continue

        room = db.query(Room).filter(Room.room_code == room_code).first()
        if not room:
            skipped_count += 1
            errors.append(f"第{index}行 room_code 不存在: {room_code}")
            continue

        quantity_text = row.get("quantity", "1").strip()
        try:
            quantity = max(1, int(float(quantity_text or "1")))
        except ValueError:
            skipped_count += 1
            errors.append(f"第{index}行 quantity 非法: {quantity_text}")
            continue

        asset = db.query(Asset).filter(Asset.asset_code == asset_code).first()
        if asset:
            asset.asset_name = asset_name
            asset.asset_category = row.get("asset_category", asset.asset_category) or asset.asset_category
            asset.room_id = room.id
            asset.quantity = quantity
            asset.status = row.get("status", asset.status) or asset.status
            asset.manufacturer = row.get("manufacturer") or None
            asset.model = row.get("model") or None
            asset.note = row.get("note") or None
            updated_count += 1
        else:
            db.add(
                Asset(
                    asset_code=asset_code,
                    asset_name=asset_name,
                    asset_category=row.get("asset_category", "weak_current") or "weak_current",
                    room_id=room.id,
                    quantity=quantity,
                    status=row.get("status", "in_use") or "in_use",
                    manufacturer=row.get("manufacturer") or None,
                    model=row.get("model") or None,
                    note=row.get("note") or None,
                )
            )
            created_count += 1

    db.commit()
    return ImportSummary(
        total_rows=len(rows),
        created_count=created_count,
        updated_count=updated_count,
        skipped_count=skipped_count,
        errors=errors,
    )
